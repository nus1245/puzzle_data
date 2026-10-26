"""측정 기반 레이아웃 엔진 (python-pptx).

원칙
  1. 모든 텍스트는 Pillow 로 실측 → 어절 단위 균형 줄 계획 → 명시적 줄바꿈(<a:br/>)으로 삽입.
  2. 박스 높이는 (줄 수 × 고정 행간)으로 계산. 카드는 위→아래 흐름으로 좌표를 산출해 겹침이 구조적으로 불가능.
  3. 넘치면 폰트를 줄이지 않고 LayoutError(빌드 실패). 문구를 줄이거나 레이아웃을 바꾼다.
  4. 숫자는 Facts 에서만 주입하고, 슬라이드 텍스트의 모든 숫자 리터럴을 facts 와 대조한다.
"""
import json
import os
import re
import uuid
from typing import List, Optional, Sequence

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FONT_DIR = r"C:\Windows\Fonts"

# ------------------------------------------------------------------ 토큰
SLIDE_W, SLIDE_H = 13.333, 7.5
MX = 0.55                                  # 좌우 여백
CW = SLIDE_W - 2 * MX                       # 콘텐츠 폭 12.233
BG, CARD, BORDER = "F8FAFC", "FFFFFF", "E2E8F0"
INK, INK2, MUTED = "0F172A", "334155", "64748B"
ACC, RED, ORG, GRN = "0284C7", "DC2626", "EA580C", "059669"
ACC_BG, RED_BG, ORG_BG, GRN_BG, SLATE_BG = "E0F2FE", "FEF2F2", "FFF7ED", "ECFDF5", "F1F5F9"
K, M = "맑은 고딕", "Consolas"             # 본문/제목, 태그·숫자
FIT = 0.965                                # 줄 폭 안전율
LH_BODY, LH_MONO, LH_TITLE, LH_BIG = 1.40, 1.32, 1.28, 1.12


class LayoutError(Exception):
    pass


# ------------------------------------------------------------------ 폰트 측정
_FILES = {("k", 0): "malgun.ttf", ("k", 1): "malgunbd.ttf", ("m", 0): "consola.ttf", ("m", 1): "consolab.ttf"}
_CACHE = {}


def _font(kind, bold, size):
    key = (kind, bool(bold), round(size * 20))
    if key not in _CACHE:
        _CACHE[key] = ImageFont.truetype(os.path.join(FONT_DIR, _FILES[(kind, int(bool(bold)))]), int(round(size * 20)))
    return _CACHE[key]


def _is_hangul(ch):
    o = ord(ch)
    return 0xAC00 <= o <= 0xD7A3 or 0x3130 <= o <= 0x318F or 0x1100 <= o <= 0x11FF


def char_w(ch, kind, bold, size):
    """한 글자의 폭(inch). Consolas 런에서도 한글·기호는 맑은 고딕으로 렌더되므로 그 폰트로 측정."""
    k = kind
    if kind == "m" and (ord(ch) > 127):
        k = "k"
    return _font(k, bold, size).getlength(ch) / 1440.0


def text_w(text, kind, bold, size):
    return sum(char_w(c, kind, bold, size) for c in text)


# ------------------------------------------------------------------ 런/문단
class R:
    def __init__(self, text, size=11.5, bold=False, color=INK, font="k"):
        self.text, self.size, self.bold, self.color, self.font = text, size, bold, color, font

    def style(self):
        return (self.size, self.bold, self.color, self.font)


class P:
    """문단. runs: [R]. bullet=True 이면 매달린 들여쓰기."""

    def __init__(self, runs, align="l", bullet=False, lh=None, before=0.0, balance=True, bullet_color=ACC):
        self.runs = runs if isinstance(runs, list) else [runs]
        self.align, self.bullet, self.before, self.balance, self.bullet_color = align, bullet, before, balance, bullet_color
        sz = max(r.size for r in self.runs)
        mono = all(r.font == "m" for r in self.runs)
        self.lh = lh or (LH_MONO if mono else LH_BODY)
        self.size = sz

    @property
    def pitch_pt(self):
        return round(self.size * self.lh, 1)


def T(text, size=11.5, bold=False, color=INK, font="k", **kw):
    return P([R(text, size, bold, color, font)], **kw)


def tri_runs(text, size, bold, color, font="k"):
    """▼▲ 글리프는 본문보다 크게 렌더되므로 62% 크기로 분리한다."""
    runs = []
    for part in re.split(r"([▼▲])", text):
        if not part:
            continue
        runs.append(R(part, size * 0.62, bold, color, "k") if part in "▼▲" else R(part, size, bold, color, font))
    return runs


def rich(text, size=11.5, color=INK, accent=ACC, font="k", bold=False, **kw):
    """'**강조**' 구간을 굵은 accent 색으로. 나머지는 기본 스타일."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    runs = []
    for s in parts:
        if not s:
            continue
        if s.startswith("**"):
            runs.append(R(s[2:-2], size, True, accent, font))
        else:
            runs.append(R(s, size, bold, color, font))
    return P(runs, **kw)


# ------------------------------------------------------------------ 줄 계획
def _words(p: P):
    """문단을 단어(공백으로 분리) 리스트로. 단어 = [(text, run_idx)] 구간 목록."""
    words, cur = [], []
    for ri, r in enumerate(p.runs):
        buf = ""
        for ch in r.text:
            if ch == " ":
                if buf:
                    cur.append((buf, ri))
                    buf = ""
                if cur:
                    words.append(cur)
                    cur = []
            else:
                buf += ch
        if buf:
            cur.append((buf, ri))
    if cur:
        words.append(cur)
    merged = []
    for w in words:
        if "".join(t for t, _ in w) in ("·", "•", "|", "—") and merged:
            merged[-1] = merged[-1] + [(" ", merged[-1][-1][1])] + w
        else:
            merged.append(w)
    return merged


def _wlen(word, p):
    return sum(text_w(t, p.runs[ri].font, p.runs[ri].bold, p.runs[ri].size) for t, ri in word)


def _spw(p, ri=0):
    r = p.runs[ri]
    return text_w(" ", r.font, r.bold, r.size)


def plan_lines(p: P, width: float):
    """균형 줄 계획. 반환: [[word,...], ...]"""
    words = _words(p)
    if not words:
        return [[]]
    lim = width * FIT
    wl = [_wlen(w, p) for w in words]
    sp = _spw(p, words[0][0][1])
    for w, l in zip(words, wl):
        if l > lim:
            raise LayoutError(f"단어가 줄 폭({lim:.2f}in)보다 김: '{''.join(t for t, _ in w)}' ({l:.2f}in)")
    # 탐욕 줄 수
    lines, cur, cw_ = [], [], 0.0
    for w, l in zip(words, wl):
        add = l if not cur else sp + l
        if cur and cw_ + add > lim:
            lines.append(cur)
            cur, cw_ = [w], l
        else:
            cur.append(w)
            cw_ += add
    lines.append(cur)
    n = len(lines)
    if n == 1 or not p.balance:
        return lines
    # 정확히 n줄로 나누되 들쭉날쭉(제곱 여백)을 최소화하는 DP
    m = len(words)
    pre = [0.0]
    for l in wl:
        pre.append(pre[-1] + l)

    def seg(i, j):               # words[i:j] 한 줄의 폭
        return pre[j] - pre[i] + sp * (j - i - 1)

    INF = 1e18
    dp = [[INF] * (m + 1) for _ in range(n + 1)]
    bk = [[0] * (m + 1) for _ in range(n + 1)]
    dp[0][0] = 0.0
    for j in range(1, n + 1):
        for i in range(j, m + 1):
            for k in range(j - 1, i):
                if dp[j - 1][k] >= INF:
                    continue
                w_ = seg(k, i)
                if w_ > lim:
                    continue
                slack = lim - w_
                cost = dp[j - 1][k] + (slack ** 2) * (0.5 if j == n else 1.0)
                if cost < dp[j][i]:
                    dp[j][i], bk[j][i] = cost, k
    if dp[n][m] >= INF:
        return lines
    out, i = [], m
    for j in range(n, 0, -1):
        k = bk[j][i]
        out.append(words[k:i])
        i = k
    return out[::-1]


def para_height_in(p: P, width: float):
    marl = 0.18 if p.bullet else 0.0
    n = len(plan_lines(p, width - marl))
    return (n * p.pitch_pt + p.before) / 72.0, n


# ------------------------------------------------------------------ XML 빌더
_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def _rpr(r: R, tag="a:rPr"):
    e = etree.Element(qn(tag))
    e.set("lang", "ko-KR")
    e.set("sz", str(int(round(r.size * 100))))
    e.set("b", "1" if r.bold else "0")
    e.set("dirty", "0")
    sf = etree.SubElement(e, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", r.color)
    latin = K if r.font == "k" else M
    etree.SubElement(e, qn("a:latin")).set("typeface", latin)
    etree.SubElement(e, qn("a:ea")).set("typeface", K)
    etree.SubElement(e, qn("a:cs")).set("typeface", latin)
    return e


def _build_para(p: P, width: float):
    marl = 0.18 if p.bullet else 0.0
    lines = plan_lines(p, width - marl)
    ap = etree.Element(qn("a:p"))
    ppr = etree.SubElement(ap, qn("a:pPr"))
    ppr.set("algn", {"l": "l", "c": "ctr", "r": "r"}[p.align])
    if p.bullet:
        ppr.set("marL", str(int(marl * 914400)))
        ppr.set("indent", str(-int(marl * 914400)))
    ls = etree.SubElement(ppr, qn("a:lnSpc"))
    etree.SubElement(ls, qn("a:spcPts")).set("val", str(int(round(p.pitch_pt * 100))))
    sb = etree.SubElement(ppr, qn("a:spcBef"))
    etree.SubElement(sb, qn("a:spcPts")).set("val", str(int(round(p.before * 100))))
    if p.bullet:
        bc = etree.SubElement(ppr, qn("a:buClr"))
        etree.SubElement(bc, qn("a:srgbClr")).set("val", p.bullet_color)
        etree.SubElement(ppr, qn("a:buFont")).set("typeface", K)
        etree.SubElement(ppr, qn("a:buChar")).set("char", "•")
    else:
        etree.SubElement(ppr, qn("a:buNone"))
    for li, line in enumerate(lines):
        if li > 0:
            br = etree.SubElement(ap, qn("a:br"))
            br.append(_rpr(p.runs[line[0][0][1]] if line and line[0] else p.runs[0]))
        # 줄 안에서 같은 스타일의 연속 구간을 합친다
        segs = []                              # [(run_idx, text)]
        for wi, w in enumerate(line):
            if wi > 0:
                segs.append((segs[-1][0], " "))
            for t, ri in w:
                segs.append((ri, t))
        merged = []
        for ri, t in segs:
            if merged and merged[-1][0] == ri:
                merged[-1][1] += t
            else:
                merged.append([ri, t])
        for ri, t in merged:
            r = etree.SubElement(ap, qn("a:r"))
            r.append(_rpr(p.runs[ri]))
            etree.SubElement(r, qn("a:t")).text = t
    return ap, len(lines)


def _hex(c):
    return RGBColor.from_string(c)


# ------------------------------------------------------------------ Facts
class Facts:
    def __init__(self, path=None):
        path = path or os.path.join(ROOT, "scripts", "deck", "facts.json")
        self.d = json.load(open(path, encoding="utf-8"))["facts"]
        self.used = set()
        self.log = {}                         # 슬라이드 번호 -> {key}
        self.cur = 0

    def __call__(self, key):
        self.used.add(key)
        self.log.setdefault(self.cur, set()).add(key)
        return self.d[key]["s"]

    def allowed_numbers(self):
        out = set()
        for k, v in self.d.items():
            for tok in re.findall(r"\d[\d,]*\.?\d*", v["s"]):
                out.add(tok.replace(",", "").rstrip("."))
        return out


def _strip_style(shape):
    """python-pptx 커넥터의 <p:style>(테마 effectRef=그림자)을 제거해 1px 선만 남긴다."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


# ------------------------------------------------------------------ 문서/슬라이드
class Doc:
    def __init__(self, facts: Facts):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
        self.facts = facts
        self.slides: List["Slide"] = []
        self.links = []                       # (shape, target_slide_no)

    def new_slide(self, no, section):
        s = Slide(self, no, section)
        self.slides.append(s)
        self.facts.cur = no
        return s


class Slide:
    def __init__(self, doc: Doc, no: int, section: str):
        self.doc, self.no, self.section = doc, no, section
        self.s = doc.prs.slides.add_slide(doc.prs.slide_layouts[6])
        bg = self.s.background.fill
        bg.solid()
        bg.fore_color.rgb = _hex(BG)
        self.reg = []                          # 지오메트리 레지스트리
        self.texts = []                        # (text, structural) 숫자 검사용
        self._id = 0

    # ---- 레지스트리
    def _reg(self, kind, x, y, w, h, parent, name):
        self._id += 1
        self.reg.append(dict(id=self._id, kind=kind, x=x, y=y, w=w, h=h, parent=parent, name=name))
        return self._id

    # ---- 도형
    def rect(self, x, y, w, h, fill=CARD, line=BORDER, lw=0.75, parent=0, name="card", kind="box"):
        sh = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.shadow.inherit = False
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = _hex(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = _hex(line)
            sh.line.width = Pt(lw)
        else:
            sh.line.fill.background()
        sh.name = name
        return sh, self._reg(kind, x, y, w, h, parent, name)

    def hline(self, x, y, w, color=BORDER, lw=0.75, parent=0, name="rule"):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
        c.line.color.rgb = _hex(color)
        c.line.width = Pt(lw)
        c.name = name
        _strip_style(c)
        self._reg("line", x, y, w, 0.0, parent, name)
        return c

    def vline(self, x, y, h, color=BORDER, lw=0.75, parent=0, name="rule"):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x), Inches(y + h))
        c.line.color.rgb = _hex(color)
        c.line.width = Pt(lw)
        c.name = name
        _strip_style(c)
        self._reg("line", x, y, 0.0, h, parent, name)
        return c

    # ---- 텍스트
    def text(self, x, y, w, paras, parent=0, name="txt", anchor="t", h=None, structural=False, role="body", max_h=None, max_lines=None):
        paras = paras if isinstance(paras, (list, tuple)) else [paras]
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.3))
        txb = tb.text_frame._txBody
        bp = txb.find(qn("a:bodyPr"))
        for k, v in dict(wrap="square", lIns="0", tIns="0", rIns="0", bIns="0", anchor={"t": "t", "m": "ctr", "b": "b"}[anchor]).items():
            bp.set(k, v)
        for ch in list(bp):
            bp.remove(ch)
        etree.SubElement(bp, qn("a:noAutofit"))
        for p_ in txb.findall(qn("a:p")):
            txb.remove(p_)
        total_pt, nlines = 0.0, 0
        for pi, p in enumerate(paras):
            ap, n = _build_para(p, w)
            txb.append(ap)
            total_pt += n * p.pitch_pt + (p.before if pi > 0 else 0.0)
            nlines += n
            self.texts.append(("".join(r.text for r in p.runs), structural))
        # 첫 문단의 before 는 무시(상단 정렬 기준)
        first = txb.findall(qn("a:p"))[0].find(qn("a:pPr")).find(qn("a:spcBef"))
        if first is not None:
            first.find(qn("a:spcPts")).set("val", "0")
        th = total_pt / 72.0
        if max_h is not None and th > max_h + 1e-6:
            raise LayoutError(f"[S{self.no}:{name}] 텍스트 높이 {th:.2f}in > 허용 {max_h:.2f}in ({nlines}줄): {paras[0].runs[0].text[:30]}")
        if max_lines is not None and nlines > max_lines:
            raise LayoutError(f"[S{self.no}:{name}] {nlines}줄 > 허용 {max_lines}줄: {paras[0].runs[0].text[:40]}")
        box_h = h if h is not None else th
        if box_h + 1e-6 < th:
            raise LayoutError(f"[S{self.no}:{name}] 박스 높이 {box_h:.2f}in < 텍스트 {th:.2f}in")
        tb.height = Inches(box_h)
        tb.name = f"{role}:L{nlines}:{name}"
        self._reg("text", x, y, w, box_h, parent, name)
        return tb, th, nlines

    def picture(self, path, x, y, w, h, parent=0, name="pic", pad=0.0, align="c"):
        """contain 배치. 종횡비 보존, (x,y,w,h) 안에서 가운데 정렬."""
        iw, ih = Image.open(path).size
        asp = iw / ih
        bw, bh = w - 2 * pad, h - 2 * pad
        if bw / bh > asp:
            ph, pw = bh, bh * asp
        else:
            pw, ph = bw, bw / asp
        px = x + pad + (bw - pw) / 2 if align == "c" else x + pad
        py = y + pad + (bh - ph) / 2
        pic = self.s.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(ph))
        pic.name = f"pic:{name}"
        self._reg("pic", px, py, pw, ph, parent, name)
        return pic, pw, ph

    def notes(self, text):
        self.s.notes_slide.notes_text_frame.text = text

    # ---- 검증
    def finish(self):
        """겹침·경계 검사. 같은 부모의 형제끼리, 또는 서로 다른 계열의 leaf 끼리 교차하면 실패."""
        byid = {r["id"]: r for r in self.reg}

        def ancestors(r):
            out, p = set(), r["parent"]
            while p:
                out.add(p)
                p = byid[p]["parent"]
            return out

        eps = 0.004
        for r in self.reg:
            if r["kind"] in ("line",):
                continue
            if r["x"] < 0.4 - eps or r["x"] + r["w"] > SLIDE_W - 0.4 + eps or r["y"] < 0.3 - eps or r["y"] + r["h"] > SLIDE_H - 0.2 + eps:
                raise LayoutError(f"[S{self.no}] 안전 영역 이탈: {r['name']} ({r['x']:.2f},{r['y']:.2f},{r['w']:.2f},{r['h']:.2f})")
        items = [r for r in self.reg if r["kind"] != "line"]
        for i, a in enumerate(items):
            for b in items[i + 1:]:
                if a["id"] in ancestors(b) or b["id"] in ancestors(a):
                    # 부모-자식: 자식이 부모 안에 있어야 한다
                    par, ch = (a, b) if a["id"] in ancestors(b) else (b, a)
                    if ch["x"] < par["x"] - eps or ch["y"] < par["y"] - eps or ch["x"] + ch["w"] > par["x"] + par["w"] + eps or ch["y"] + ch["h"] > par["y"] + par["h"] + eps:
                        raise LayoutError(f"[S{self.no}] 자식이 부모 밖: {ch['name']} ⊄ {par['name']}")
                    continue
                ox = min(a["x"] + a["w"], b["x"] + b["w"]) - max(a["x"], b["x"])
                oy = min(a["y"] + a["h"], b["y"] + b["h"]) - max(a["y"], b["y"])
                if ox > eps and oy > eps:
                    raise LayoutError(f"[S{self.no}] 겹침: {a['name']} × {b['name']} ({ox:.2f}×{oy:.2f}in)")


# ------------------------------------------------------------------ 숫자 리터럴 검사
NUM = re.compile(r"\d[\d,]*\.?\d*")


def check_numbers(slide: Slide, facts: Facts, extra_allow: Sequence[str] = ()):
    allow = facts.allowed_numbers() | {a.replace(",", "") for a in extra_allow}
    bad = []
    for text, structural in slide.texts:
        if structural:
            continue
        for tok in NUM.findall(text):
            t = tok.replace(",", "").rstrip(".")
            if len(t) == 1:                   # 한 자리 정수(5대, 3종, E-1 등)는 개수 표기로 허용
                continue
            if t not in allow:
                bad.append((t, text[:40]))
    return bad
