"""PowerPoint 실측 결과(audit.json)로 게이트 A–F 를 판정한다.

  A 제목 1줄 / B 텍스트 경계 넘침 / C 줄 수 일치 / F 최소 글자 크기
실행: python -X utf8 scripts/deck/audit.py <pptx> <audit.json>
종료코드: 위반 0 이면 0
"""
import json
import re
import sys

from pptx import Presentation

TOL_PT = 1.5
MIN_PT = {"body": 11.0}          # 그 외 역할(label/pill/tag/foot/page/brand/sub/title)은 9pt 이상


def main(pptx, js):
    rows = json.load(open(js, encoding="utf-8-sig"))
    if isinstance(rows, dict):
        rows = [rows]
    bad = []
    n_text = 0
    for r in rows:
        n_text += 1
        m = re.match(r"(\w+):L(\d+):(.*)", r["name"])
        if not m:
            continue
        role, planned, nm = m.group(1), int(m.group(2)), m.group(3)
        tag = f"S{r['slide']:02d} {r['name']}"
        if r["lines"] != planned:
            bad.append(f"C 줄 불일치 {tag}: 계획 {planned} / 실제 {r['lines']}")
        if r["boundH"] > r["height"] + TOL_PT:
            bad.append(f"B 높이 넘침 {tag}: 텍스트 {r['boundH']}pt > 박스 {r['height']}pt")
        if r["boundW"] > r["width"] + TOL_PT:
            bad.append(f"B 폭 넘침 {tag}: 텍스트 {r['boundW']}pt > 박스 {r['width']}pt")
        if role == "title" and r["lines"] != 1 and r["slide"] != 1:
            bad.append(f"A 제목 {r['lines']}줄 {tag}")
    prs = Presentation(pptx)
    for si, sl in enumerate(prs.slides, 1):
        for sh in sl.shapes:
            if not sh.has_text_frame:
                continue
            role = sh.name.split(":")[0]
            lim = MIN_PT.get(role, 9.0)
            for p in sh.text_frame.paragraphs:
                for rn in p.runs:
                    if rn.font.size is not None and rn.font.size.pt + 1e-6 < lim:
                        bad.append(f"F 글자 {rn.font.size.pt}pt < {lim}pt S{si:02d} {sh.name}: {rn.text[:20]}")
    print(f"텍스트 도형 {n_text}개 감사, 위반 {len(bad)}건")
    for b in bad[:80]:
        print("  -", b)
    return 0 if not bad else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1], sys.argv[2]))
